import pytest
from openpyxl import load_workbook

from core.tools import TOOLS
from core.tools.base import ToolError
from core.tools.reports import sheet_name

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def test_sheet_name_sanitized():
    used: set[str] = set()
    assert sheet_name("a[b]:c*d?e/f\\g", used) == "abcdefg"
    assert len(sheet_name("x" * 50, used)) == 31
    assert sheet_name("abcdefg", used) == "abcdefg 2"
    assert sheet_name("[]", used) == "Sheet"


@pytest.mark.parametrize("tool,args", [
    ("make_excel", {"title": "t", "sheets": []}),
    ("make_excel", {"title": "t", "sheets": [{"name": "a", "columns": ["x"], "rows": [[1]]}] * 21}),
    ("make_excel", {"title": "t", "sheets": [{"name": "a", "columns": ["x"], "rows": [[1]] * 5001}]}),
    ("make_excel", {"title": "t", "sheets": [{"name": "a", "columns": ["x"], "rows": [[[1]]]}]}),
    ("make_excel", {"title": "", "sheets": [{"name": "a", "columns": ["x"], "rows": []}]}),
    ("make_excel", {"title": "t", "sheets": [{"name": "a", "columns": ["x"], "rows": [], "z": 1}]}),
    ("make_pdf", {"title": "t", "sections": [{"heading": "h", "paragraphs": "no"}]}),
    ("make_pdf", {"title": "t", "sections": [], "extra": 1}),
    ("make_chart", {"kind": "radar", "labels": ["a"], "values": [1], "title": "t"}),
    ("make_chart", {"kind": "bar", "labels": ["a", "b"], "values": [1], "title": "t"}),
    ("make_chart", {"kind": "bar", "labels": ["a"], "values": ["1"], "title": "t"}),
    ("make_chart", {"kind": "pie", "labels": ["a"], "values": [-1], "title": "t"}),
    ("docker_report", {"x": 1}),
])
def test_validation_rejects(tool, args):
    with pytest.raises(ToolError):
        TOOLS[tool].validate(args)


async def test_make_excel(ctx):
    args = TOOLS["make_excel"].validate({
        "title": "Ventas Q1",
        "sheets": [{"name": "Data/1", "columns": ["name", "=amount"],
                    "rows": [["al\x00pha", 10], ["=1+1", 2.5]]}],
    })
    r = await TOOLS["make_excel"].run(args, ctx)
    f = r.files[0]
    assert f.mime == XLSX and f.path.parent == ctx.out_dir and f.path.suffix == ".xlsx"
    ws = load_workbook(f.path)["Data1"]
    assert ws["A1"].font.bold and ws["B1"].font.bold
    assert ws.auto_filter.ref == "A1:B3"
    assert ws["A2"].value == "alpha" and ws["B3"].value == 2.5
    assert ws["A3"].data_type == "s"
    assert ws["B1"].data_type == "s"
    assert ws.column_dimensions["A"].width >= len("alpha")
    r2 = await TOOLS["make_excel"].run(args, ctx)
    assert r2.files[0].path != f.path


async def test_make_pdf(ctx):
    args = TOOLS["make_pdf"].validate({
        "title": "Report <b>&",
        "sections": [{"heading": "H <i>", "paragraphs": ["a < b & <unclosed>"],
                      "table": {"columns": ["c1"], "rows": [["<x>"], [None]]}}],
    })
    r = await TOOLS["make_pdf"].run(args, ctx)
    f = r.files[0]
    assert f.mime == "application/pdf"
    data = f.path.read_bytes()
    assert data.startswith(b"%PDF") and len(data) > 0


async def test_make_chart(ctx):
    for kind in ("bar", "line", "pie"):
        args = TOOLS["make_chart"].validate({"kind": kind, "labels": ["a", "b"], "values": [1, 2], "title": kind})
        r = await TOOLS["make_chart"].run(args, ctx)
        assert r.files[0].mime == "image/png"
        assert r.files[0].path.read_bytes().startswith(b"\x89PNG")


async def test_docker_report(ctx):
    r = await TOOLS["docker_report"].run({}, ctx)
    assert sorted(f.path.suffix for f in r.files) == [".pdf", ".xlsx"]
    assert all(f.path.stat().st_size > 0 for f in r.files)
    assert len(list(ctx.out_dir.iterdir())) == 2
    xlsx = next(f.path for f in r.files if f.path.suffix == ".xlsx")
    ws = load_workbook(xlsx).active
    assert [c.value for c in ws["A"]] == ["Container", "db", "old", "web"]
