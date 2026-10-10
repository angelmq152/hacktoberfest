import asyncio
import re
import tempfile
from datetime import datetime
from pathlib import Path
from uuid import uuid4
from xml.sax.saxutils import escape

import matplotlib
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from core.tools import docker_tools
from core.tools.base import ToolContext, ToolError, ToolResult, ToolSpec
from shared.contracts import OutFile

matplotlib.use("Agg")
from matplotlib.figure import Figure  # noqa: E402

MAX_SHEETS = 20
MAX_ROWS = 5000
MAX_COLUMNS = 50
MAX_SECTIONS = 30
MAX_PARAGRAPHS = 50
MAX_PDF_ROWS = 300
MAX_POINTS = 100
MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MIME_PDF = "application/pdf"
MIME_PNG = "image/png"
_BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")


def _check_args(args, allowed: set[str]) -> dict:
    if not isinstance(args, dict):
        raise ToolError("args must be an object")
    extra = set(args) - allowed
    if extra:
        raise ToolError(f"unexpected args: {', '.join(sorted(extra))}")
    return args


def _title(value) -> str:
    if not isinstance(value, str) or not value.strip():
        raise ToolError("'title' must be a non-empty string")
    return value.strip()[:200]


def _cell(value):
    if value is None or isinstance(value, (bool, int, float)):
        return value
    if isinstance(value, str):
        return value
    raise ToolError("cells must be string, number, boolean or null")


def _table(columns, rows, max_rows: int, where: str) -> tuple[list[str], list[list]]:
    if not isinstance(columns, list) or not columns or len(columns) > MAX_COLUMNS:
        raise ToolError(f"{where}: 'columns' must be a list of 1-{MAX_COLUMNS} strings")
    if not all(isinstance(c, str) for c in columns):
        raise ToolError(f"{where}: 'columns' must be strings")
    if not isinstance(rows, list) or len(rows) > max_rows:
        raise ToolError(f"{where}: 'rows' must be a list with at most {max_rows} rows")
    clean = []
    for r in rows:
        if not isinstance(r, list) or len(r) > len(columns):
            raise ToolError(f"{where}: each row must be a list with at most {len(columns)} cells")
        clean.append([_cell(v) for v in r])
    return columns, clean


def _out_path(ctx: ToolContext, title: str, ext: str) -> Path:
    ctx.out_dir.mkdir(parents=True, exist_ok=True)
    slug = re.sub(r"[^a-z0-9]+", "-", title.lower()).strip("-")[:40] or "report"
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    return ctx.out_dir / f"{slug}-{stamp}-{uuid4().hex[:6]}.{ext}"


def sheet_name(name: str, used: set[str]) -> str:
    base = _BAD_SHEET_CHARS.sub("", name).strip("'").strip()[:31] or "Sheet"
    out, i = base, 2
    while out.lower() in used:
        suffix = f" {i}"
        out = base[: 31 - len(suffix)] + suffix
        i += 1
    used.add(out.lower())
    return out


# make_excel
def _v_excel(args: dict) -> dict:
    args = _check_args(args, {"title", "sheets"})
    sheets = args.get("sheets")
    if not isinstance(sheets, list) or not 1 <= len(sheets) <= MAX_SHEETS:
        raise ToolError(f"'sheets' must be a list of 1-{MAX_SHEETS} sheets")
    clean = []
    for i, s in enumerate(sheets):
        s = _check_args(s, {"name", "columns", "rows"})
        if not isinstance(s.get("name"), str):
            raise ToolError(f"sheet {i}: 'name' must be a string")
        cols, rows = _table(s.get("columns"), s.get("rows", []), MAX_ROWS, f"sheet {i}")
        clean.append({"name": s["name"], "columns": cols, "rows": rows})
    return {"title": _title(args.get("title")), "sheets": clean}


def _clean_cell(value):
    return ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value


def build_excel(path: Path, sheets: list[dict]) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    for s in sheets:
        ws = wb.create_sheet(sheet_name(s["name"], used))
        ws.append([_clean_cell(c) for c in s["columns"]])
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for r in s["rows"]:
            ws.append([_clean_cell(c) for c in r])
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.data_type = "s"
        for i, col in enumerate(s["columns"], start=1):
            width = max([len(str(col))] + [len(str(r[i - 1])) for r in s["rows"] if i <= len(r) and r[i - 1] is not None])
            ws.column_dimensions[get_column_letter(i)].width = min(width + 2, 60)
        ws.auto_filter.ref = ws.dimensions
    wb.save(path)


async def _run_excel(args: dict, ctx: ToolContext) -> ToolResult:
    path = _out_path(ctx, args["title"], "xlsx")
    await asyncio.to_thread(build_excel, path, args["sheets"])
    return ToolResult(data=f"Excel '{path.name}' created", files=[OutFile(path, MIME_XLSX, args["title"])])


# make_pdf
def _v_pdf(args: dict) -> dict:
    args = _check_args(args, {"title", "sections"})
    sections = args.get("sections")
    if not isinstance(sections, list) or not 1 <= len(sections) <= MAX_SECTIONS:
        raise ToolError(f"'sections' must be a list of 1-{MAX_SECTIONS} sections")
    clean = []
    for i, s in enumerate(sections):
        s = _check_args(s, {"heading", "paragraphs", "table"})
        heading = s.get("heading", "")
        paragraphs = s.get("paragraphs", [])
        if not isinstance(heading, str):
            raise ToolError(f"section {i}: 'heading' must be a string")
        if (not isinstance(paragraphs, list) or len(paragraphs) > MAX_PARAGRAPHS
                or not all(isinstance(p, str) for p in paragraphs)):
            raise ToolError(f"section {i}: 'paragraphs' must be a list of at most {MAX_PARAGRAPHS} strings")
        sec = {"heading": heading, "paragraphs": paragraphs, "table": None}
        if s.get("table") is not None:
            t = _check_args(s["table"], {"columns", "rows"})
            cols, rows = _table(t.get("columns"), t.get("rows", []), MAX_PDF_ROWS, f"section {i} table")
            sec["table"] = {"columns": cols, "rows": rows}
        clean.append(sec)
    return {"title": _title(args.get("title")), "sections": clean}


def build_pdf(path: Path, title: str, sections: list[dict], images: list[Path] | None = None) -> None:
    styles = getSampleStyleSheet()
    cell_style = styles["BodyText"]
    story = [Paragraph(escape(title), styles["Title"]), Spacer(1, 0.4 * cm)]
    for s in sections:
        if s["heading"]:
            story.append(Paragraph(escape(s["heading"]), styles["Heading2"]))
        story += [Paragraph(escape(p), styles["BodyText"]) for p in s["paragraphs"]]
        if s["table"]:
            data = [[Paragraph(f"<b>{escape(c)}</b>", cell_style) for c in s["table"]["columns"]]]
            data += [[Paragraph(escape("" if v is None else str(v)), cell_style) for v in r] for r in s["table"]["rows"]]
            t = Table(data, repeatRows=1)
            t.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]))
            story.append(t)
        story.append(Spacer(1, 0.4 * cm))
    for img in images or []:
        story.append(Image(str(img), width=15 * cm, height=9 * cm))
    SimpleDocTemplate(str(path), pagesize=A4, title=title[:100]).build(story)


async def _run_pdf(args: dict, ctx: ToolContext) -> ToolResult:
    path = _out_path(ctx, args["title"], "pdf")
    await asyncio.to_thread(build_pdf, path, args["title"], args["sections"])
    return ToolResult(data=f"PDF '{path.name}' created", files=[OutFile(path, MIME_PDF, args["title"])])


# make_chart
def _v_chart(args: dict) -> dict:
    args = _check_args(args, {"kind", "labels", "values", "title"})
    kind = args.get("kind")
    if kind not in ("bar", "line", "pie"):
        raise ToolError("'kind' must be one of bar, line, pie")
    labels, values = args.get("labels"), args.get("values")
    if not isinstance(labels, list) or not all(isinstance(x, str) for x in labels):
        raise ToolError("'labels' must be a list of strings")
    if (not isinstance(values, list)
            or not all(isinstance(x, (int, float)) and not isinstance(x, bool) for x in values)):
        raise ToolError("'values' must be a list of numbers")
    if len(labels) != len(values):
        raise ToolError("'labels' and 'values' must have the same length")
    if not 1 <= len(values) <= MAX_POINTS:
        raise ToolError(f"chart needs 1-{MAX_POINTS} points")
    if kind == "pie" and any(v < 0 for v in values):
        raise ToolError("pie charts do not accept negative values")
    return {"kind": kind, "labels": labels, "values": values, "title": _title(args.get("title"))}


def render_chart(path: Path, kind: str, labels: list[str], values: list[float], title: str) -> None:
    fig = Figure(figsize=(8, 5))
    ax = fig.subplots()
    if kind == "bar":
        ax.bar(labels, values)
        ax.tick_params(axis="x", rotation=45)
    elif kind == "line":
        ax.plot(labels, values, marker="o")
        ax.tick_params(axis="x", rotation=45)
    else:
        ax.pie(values, labels=labels, autopct="%1.1f%%")
    ax.set_title(title)
    fig.tight_layout()
    fig.savefig(path, format="png", dpi=100)


async def _run_chart(args: dict, ctx: ToolContext) -> ToolResult:
    path = _out_path(ctx, args["title"], "png")
    await asyncio.to_thread(render_chart, path, args["kind"], args["labels"], args["values"], args["title"])
    return ToolResult(data=f"Chart '{path.name}' created", files=[OutFile(path, MIME_PNG, args["title"])])


# docker_report
def _v_report(args: dict) -> dict:
    _check_args(args, set())
    return {}


def _build_docker_report(ctx: ToolContext) -> list[OutFile]:
    rows = docker_tools.collect_rows()
    title = "Docker report"
    columns = ["Container", "State", "CPU %", "RAM MB", "Uptime", "Restarts"]
    table = [[r["name"], r["status"], round(r["cpu_percent"], 1), round(r["mem_mb"], 1), r["uptime"], r["restarts"]]
             for r in rows]
    xlsx = _out_path(ctx, title, "xlsx")
    pdf = _out_path(ctx, title, "pdf")
    build_excel(xlsx, [{"name": "Containers", "columns": columns, "rows": table}])
    running = sum(1 for r in rows if r["status"] == "running")
    sections = [{
        "heading": "Managed containers",
        "paragraphs": [f"{len(rows)} containers, {running} running. Generated {datetime.now():%Y-%m-%d %H:%M}."],
        "table": {"columns": columns, "rows": table},
    }]
    with tempfile.TemporaryDirectory() as tmp:
        images = []
        if rows:
            img = Path(tmp) / "ram.png"
            render_chart(img, "bar", [r["name"] for r in rows], [r["mem_mb"] for r in rows], "RAM usage (MB)")
            images.append(img)
        build_pdf(pdf, title, sections, images)
    return [OutFile(xlsx, MIME_XLSX, "Docker report (Excel)"), OutFile(pdf, MIME_PDF, "Docker report (PDF)")]


async def _run_report(args: dict, ctx: ToolContext) -> ToolResult:
    files = await asyncio.to_thread(_build_docker_report, ctx)
    return ToolResult(data="Docker report created (Excel + PDF)", files=files)


REPORT_TOOLS: dict[str, ToolSpec] = {
    t.name: t
    for t in [
        ToolSpec("make_excel", "Create an Excel file with bold header, autofilter and fitted columns.",
                 {"title": "file title", "sheets": "list of {name, columns: [str], rows: [[cell]]}"},
                 False, _v_excel, _run_excel),
        ToolSpec("make_pdf", "Create a PDF report.",
                 {"title": "document title",
                  "sections": "list of {heading, paragraphs: [str], table?: {columns, rows}}"},
                 False, _v_pdf, _run_pdf),
        ToolSpec("make_chart", "Create a PNG chart.",
                 {"kind": "bar | line | pie", "labels": "list of strings", "values": "list of numbers, same length",
                  "title": "chart title"}, False, _v_chart, _run_chart),
        ToolSpec("docker_report", "Snapshot of all managed containers as Excel + PDF with a RAM chart.",
                 {}, False, _v_report, _run_report),
    ]
}
